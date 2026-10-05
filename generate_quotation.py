#!/usr/bin/env python3
"""
generate_quotation.py
=====================
Builds a quotation as a WORD (.docx) file and a PDF, copying the layout of the
Enabled PH / Agile Technologies sample quote.

    python generate_quotation.py                       # uses quotation_config.py + products.csv
    python generate_quotation.py --number Q-26-700 --company "ACME Corp" --discounted-price 2320000
    python generate_quotation.py --help

What to edit
------------
  * quotation_config.py  -> top fields, discount, terms, letterhead, bank details
  * products.csv         -> the items (name, SRP, pictures, description, ...)
  * this file            -> only if you want to change the LAYOUT itself

Layout of this file
-------------------
  1. Imports & small constants
  2. Data loading           (CSV -> list of Item)
  3. Low-level Word helpers (rich text, borders, widths, floating picture)
  4. Page building blocks   (header/footer, info box, items table, summary, terms)
  5. PDF conversion
  6. generate()  <- one call builds everything (used by the GUI and the command line)
  7. Command line / main()
"""

import argparse
import base64
import csv
import dataclasses
import datetime
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION  # noqa: F401  (kept for easy extension)
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageChops, ImageOps

import quotation_config as cfg

# ============================================================================
# 1. SMALL CONSTANTS  (layout numbers copied from the sample quotation)
# ============================================================================
FONT = "Arial"      # the ONE font used for all text in the exported Word/PDF
SIZE = 10           # the ONE font size (pt) used for all text

PAGE_W, PAGE_H = 8.27, 11.69                 # A4, inches
HEADER_DISTANCE, FOOTER_DISTANCE = 1.0, 0.0          # HEADER_DISTANCE = gap above the logo (top of the page)
LETTERHEAD_HEIGHT_IN = 1.7                           # room for logo + contact lines + a small gap
MARGIN_LR, MARGIN_BOTTOM = 1.0, 1.0
MARGIN_TOP = HEADER_DISTANCE + LETTERHEAD_HEIGHT_IN  # where the body text starts (moves with the header)
FOOTER_BAR_HEIGHT_PT = 30                    # navy bar height

# column widths in inches
INFO_COLS    = (3.62, 3.34)                          # Quotation Number / Date ...
ITEM_COLS    = (0.68, 0.79, 3.84, 1.08, 1.18)        # Line | Qty | Item | Price | Total
SUMMARY_COLS = (0.50, 0.80, 3.75, 1.60)             # Line | Qty | Description | Price
TIMELINE_COLS = (2.95, 2.93)

LOGO_WIDTH_IN = 1.55                         # printed width of the logo (empty margins are trimmed first)

# product-picture bounding box (pictures are scaled to fit inside, keeping aspect ratio)
IMAGE_BOX_W_IN, IMAGE_BOX_H_IN = 2.3, 1.8    # EVERY product picture is fitted inside this box (aspect ratio kept)
IMAGE_ENLARGE_SMALL = True                   # True: small pictures are enlarged to the box; False: never enlarged
IMAGE_EXPORT_DPI = 200                       # pictures are down-sampled to this resolution (keeps files small)

WARNINGS = []     # non-fatal problems collected while building (e.g. missing pictures)

# markup tokens understood by add_rich()  (see quotation_config.py)
MARKUP = {"**": "bold", "~~": "italic", "__": "underline", "^^": "red", "@@": "link"}


# ============================================================================
# 2. DATA LOADING
# ============================================================================
@dataclass
class Item:
    """One line of the quotation == one row of products.csv."""
    quantity: int
    name: str
    srp: float
    optional: bool = False          # set with the Optional checkbox in the app (not from the CSV)
    uid: str = ""                   # id used by the app to pick items for extra order summaries
    summary_name: str = ""          # optional override for the ORDER SUMMARY wording
    total_override: float = None    # None = quantity x srp; a number = a custom total typed in the app
    images: list = field(default_factory=list)
    description: str = ""
    key_features: list = field(default_factory=list)
    specifications: list = field(default_factory=list)
    in_the_box: list = field(default_factory=list)

    @property
    def total(self):
        return self.total_override if self.total_override is not None else self.quantity * self.srp

    @property
    def summary_label(self):
        """Wording in ORDER SUMMARY: summary_name if given, else the first line of the item name."""
        return self.summary_name or self.name.split("\n")[0].strip()


def _split_list(value):
    """'a | b | c' -> ['a', 'b', 'c']   (empty cell -> [])"""
    return [part.strip() for part in (value or "").split("|") if part.strip()]


def _to_float(value):
    return float(str(value).replace(cfg.CURRENCY, "").replace(",", "").strip() or 0)


def load_items(csv_path):
    """Read products.csv and return a list of Item (blank rows are ignored)."""
    items = []
    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            if not (row.get("name") or "").strip():
                continue
            name = row["name"].strip().replace("\\n", "\n")   # type \n in the CSV for a line break
            items.append(Item(
                quantity=int(float(row.get("quantity") or 1)),
                name=name,
                srp=_to_float(row.get("srp")),
                summary_name=(row.get("summary_name") or "").strip(),
                images=[s.strip() for s in (row.get("images") or "").split(";") if s.strip()],
                description=(row.get("description") or "").strip(),
                key_features=_split_list(row.get("key_features")),
                specifications=_split_list(row.get("specifications")),
                in_the_box=_split_list(row.get("in_the_box")),
            ))
    if not items:
        sys.exit(f"No items found in {csv_path}")
    return items


# ----------------------------------------------------------------------------
# catalog (products.csv) helpers
# ----------------------------------------------------------------------------
def append_item_to_csv(csv_path, item):
    """Add `item` as a new row at the end of the catalog CSV (all existing columns are kept)."""
    csv_path = Path(csv_path)
    if csv_path.exists():
        with open(csv_path, newline="", encoding="utf-8-sig") as f:
            columns = csv.DictReader(f).fieldnames or []
    else:
        columns = []
    if not columns:
        columns = ["quantity", "name", "summary_name", "srp", "images", "description",
                   "key_features", "specifications", "in_the_box"]
    row = {"quantity": item.quantity, "name": item.name.replace("\n", "\\n"),
           "summary_name": item.summary_name, "srp": f"{item.srp:.2f}",
           "images": ";".join(item.images), "description": item.description,
           "key_features": " | ".join(item.key_features), "specifications": " | ".join(item.specifications),
           "in_the_box": " | ".join(item.in_the_box)}
    new_file = not csv_path.exists()
    if not new_file and csv_path.stat().st_size and not csv_path.read_bytes().endswith(b"\n"):
        with open(csv_path, "ab") as f:                       # last line had no line break: add one first
            f.write(b"\r\n")
    with open(csv_path, "w" if new_file else "a", newline="", encoding="utf-8-sig" if new_file else "utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        if new_file:
            writer.writeheader()
        writer.writerow({c: row.get(c, "") for c in columns})


def item_to_dict(item):
    return dataclasses.asdict(item)


def item_from_dict(data):
    known = {f.name for f in dataclasses.fields(Item)}
    return Item(**{k: v for k, v in data.items() if k in known})


def money(value, symbol=True):
    return f"{cfg.CURRENCY if symbol else ''}{value:,.2f}"


def totals(items, discounted_price):
    """Return (total, discounted_or_None). Optional items are ALWAYS part of the total."""
    return sum(i.total for i in items), discounted_price


# ============================================================================
# 3. LOW-LEVEL WORD HELPERS
# ============================================================================
def style_run(run, font=FONT, size=SIZE, bold=False, italic=False,
              underline=False, color=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_rich(par, text, font=FONT, size=SIZE, bold=False, color=None):
    """Add text to a paragraph, understanding the **bold** ~~italic~~ __underline__
    ^^red^^ @@link@@ markup and '\\n' line breaks."""
    state = {"bold": bold, "italic": False, "underline": False, "red": False, "link": False}
    buffer = ""

    def flush():
        nonlocal buffer
        if not buffer:
            return
        color_now = "FF0000" if state["red"] else ("0563C1" if state["link"] else color)
        for n, piece in enumerate(buffer.split("\n")):
            if n:
                par.add_run().add_break()
            if piece:
                style_run(par.add_run(piece), font, size, state["bold"], state["italic"],
                          state["underline"] or state["link"], color_now)
        buffer = ""

    i = 0
    while i < len(text):
        token = text[i:i + 2]
        if token in MARKUP:
            flush()
            state[MARKUP[token]] = not state[MARKUP[token]]
            i += 2
        else:
            buffer += text[i]
            i += 1
    flush()
    return par


def fmt(par, align=None, before=0, after=0, left=None, first=None, right=None, line=None):
    """Paragraph formatting shortcut (all sizes in points / inches)."""
    pf = par.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if line:
        pf.line_spacing = Pt(line)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if align is not None:
        par.alignment = align
    if left is not None:
        pf.left_indent = Inches(left)
    if right is not None:
        pf.right_indent = Inches(right)
    if first is not None:
        pf.first_line_indent = Inches(first)
    return par


def cell_text(cell, text, size=SIZE, bold=False, align=None, valign=WD_ALIGN_VERTICAL.TOP,
              font=FONT, color=None):
    """Write one line of (marked-up) text into a table cell."""
    par = cell.paragraphs[0]
    fmt(par, align=align)
    add_rich(par, text, font, size, bold, color)
    cell.vertical_alignment = valign
    return par


def make_table(doc, rows, widths, border=True):
    """Create a fixed-width, centred table with a grid (widths in inches)."""
    table = doc.add_table(rows=rows, cols=len(widths))
    table.style = "Table Grid" if border else None
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # Set the grid and every cell width (Word + LibreOffice both need this)
    for col, w in zip(table.columns, widths):
        col.width = Inches(w)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)
    return table


def row_height(row, inches):
    """Minimum row height (rows still grow if the content needs more room)."""
    row.height = Inches(inches)


def spacer(doc, points=6):
    return fmt(doc.add_paragraph(), before=0, after=points)


def bottom_border(par, size=8):
    """Horizontal rule under a paragraph."""
    pPr = par._p.get_or_add_pPr()
    pPr.append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{size}" w:space="1" w:color="000000"/></w:pBdr>'))


def shade_paragraph(par, hex_color):
    par._p.get_or_add_pPr().append(parse_xml(
        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>'))


def _ink_box(im):
    """Bounding box of the visible part of an RGBA image (not transparent and not white)."""
    alpha = im.split()[-1].point(lambda v: 255 if v > 10 else 0)
    not_white = ImageChops.difference(im.convert("RGB"), Image.new("RGB", im.size, "white")) \
        .convert("L").point(lambda v: 255 if v > 10 else 0)
    return ImageChops.multiply(alpha, not_white).getbbox()


def trimmed_logo(path):
    """Logo as an in-memory PNG with empty (transparent/white) margins removed, so its left edge
    lines up exactly with the text under it."""
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im).convert("RGBA")
    box = _ink_box(im)
    if box:
        im = im.crop(box)
    out = io.BytesIO()
    im.save(out, "PNG")
    out.seek(0)
    return out


def prepare_picture(path):
    """Open ANY picture (any size/format) and return (jpeg_stream, width_in, height_in) fitted
    inside IMAGE_BOX_W_IN x IMAGE_BOX_H_IN, down-sampled so the Word/PDF stays small."""
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im)
        if im.mode in ("RGBA", "LA", "P"):                       # flatten transparency onto white
            rgba = im.convert("RGBA")
            flat = Image.new("RGB", rgba.size, "white")
            flat.paste(rgba, mask=rgba.split()[-1])
            im = flat
        else:
            im = im.convert("RGB")
        w_px, h_px = im.size
        scale = min(IMAGE_BOX_W_IN / (w_px / IMAGE_EXPORT_DPI), IMAGE_BOX_H_IN / (h_px / IMAGE_EXPORT_DPI))
        if not IMAGE_ENLARGE_SMALL:
            scale = min(scale, 1.0)
        w_in, h_in = w_px / IMAGE_EXPORT_DPI * scale, h_px / IMAGE_EXPORT_DPI * scale
        target = (max(1, round(w_in * IMAGE_EXPORT_DPI)), max(1, round(h_in * IMAGE_EXPORT_DPI)))
        if target[0] < w_px:                                     # only ever shrink the pixel data
            im = im.resize(target, Image.LANCZOS)
        out = io.BytesIO()
        im.save(out, "JPEG", quality=92)
    out.seek(0)
    return out, w_in, h_in


def float_picture(run, path, width_in, x_in, y_in):
    """Insert a picture that floats in front of the text (used for the signature).
    x is measured from the column's left edge, y from the top of the paragraph."""
    shape = run.add_picture(str(path), width=Inches(width_in))
    inline = shape._inline
    emu = lambda inches: int(inches * 914400)
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="column"><wp:posOffset>{emu(x_in)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="paragraph"><wp:posOffset>{emu(y_in)}</wp:posOffset></wp:positionV>'
        f'</wp:anchor>')
    anchor.append(inline.extent)
    anchor.append(parse_xml(f'<wp:effectExtent {nsdecls("wp")} l="0" t="0" r="0" b="0"/>'))
    anchor.append(parse_xml(f'<wp:wrapNone {nsdecls("wp")}/>'))
    anchor.append(inline.docPr)
    frame = inline.find(qn("wp:cNvGraphicFramePr"))
    if frame is None:
        frame = parse_xml(f'<wp:cNvGraphicFramePr {nsdecls("wp")}/>')
    anchor.append(frame)
    anchor.append(inline.graphic)
    drawing = inline.getparent()
    drawing.remove(inline)
    drawing.append(anchor)


# ============================================================================
# 4. PAGE BUILDING BLOCKS
# ============================================================================
def setup_page(doc):
    """A4 page, margins, default font, letterhead header and footer bar."""
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(PAGE_W), Inches(PAGE_H)
    sec.left_margin = sec.right_margin = Inches(MARGIN_LR)
    sec.top_margin, sec.bottom_margin = Inches(MARGIN_TOP), Inches(MARGIN_BOTTOM)
    sec.header_distance, sec.footer_distance = Inches(HEADER_DISTANCE), Inches(FOOTER_DISTANCE)

    # ---- header: logo + contact lines (repeats on every page) ----
    header = sec.header
    par = fmt(header.paragraphs[0], left=0)           # same left edge as the lines below it
    if cfg.LOGO_IMAGE and Path(cfg.LOGO_IMAGE).exists():
        par.add_run().add_picture(trimmed_logo(cfg.LOGO_IMAGE), width=Inches(LOGO_WIDTH_IN))
    for text, underlined in cfg.HEADER_LINES:
        line = fmt(header.add_paragraph())
        style_run(line.add_run(text), underline=underlined, color=cfg.HEADER_TEXT_COLOR)

    # ---- footer: address + full-width navy bar ----
    footer = sec.footer
    addr = fmt(footer.paragraphs[0], after=6)
    style_run(addr.add_run(cfg.FOOTER_ADDRESS))
    bar = fmt(footer.add_paragraph(), left=-MARGIN_LR, right=-MARGIN_LR, line=FOOTER_BAR_HEIGHT_PT)
    shade_paragraph(bar, cfg.FOOTER_BAR_COLOR)
    style_run(bar.add_run(" "))


def add_info_box(doc, info):
    """'QUOTATION' title + the fill-in-the-blank box."""
    title = fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    add_rich(title, "QUOTATION", bold=True)

    table = make_table(doc, 4, INFO_COLS)
    for row in table.rows:
        row_height(row, 0.22)
    layout = [
        (("Quotation Number:", "quotation_number"), ("Date:", "date")),
        (("Company Name:", "company_name"), ("Attention To:", "attention_to")),
        (("Contact:", "contact"), ("Email:", "email")),
    ]
    for r, pair in enumerate(layout):
        for c, (label, key) in enumerate(pair):
            cell_text(table.cell(r, c), f"**{label}** {info[key]}".rstrip())
    last = table.cell(3, 0).merge(table.cell(3, 1))          # Address spans both columns
    cell_text(last, f"**Address:** {info['address']}".rstrip())


def add_items_table(doc, items):
    """Page 1-2: Line Item | Quantity | Item (name, pictures, details) | Price | Total."""
    spacer(doc, 14)
    table = make_table(doc, 1 + len(items), ITEM_COLS)
    row_height(table.rows[0], 0.34)

    for cell, label in zip(table.rows[0].cells, ("Line Item", "Quantity", "Item", "Price", "Total")):
        cell_text(cell, label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  valign=WD_ALIGN_VERTICAL.CENTER)

    for n, item in enumerate(items, start=1):
        row = table.rows[n].cells
        cell_text(row[0], f"{n}.0", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[1], str(item.quantity), align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_item_cell(row[2], item)
        cell_text(row[3], money(item.srp))
        cell_text(row[4], money(item.total))


def fill_item_cell(cell, item):
    """Name, pictures, description and the detail sections of one item."""
    # name (+ red "(Optional)")
    name = fmt(cell.paragraphs[0])               # default spacing = same top padding as the other cells
    add_rich(name, item.name, bold=True)
    if item.optional:
        add_rich(name, " (Optional)", bold=True, color="FF0000")

    # pictures, one per paragraph -- all fitted into the same box whatever their original size
    for filename in item.images:
        path = Path(cfg.IMAGES_DIR) / filename
        if not path.exists():
            WARNINGS.append(f"Picture not found, skipped: {path}")
            continue
        try:
            stream, w, h = prepare_picture(path)
        except Exception as exc:                                  # noqa: BLE001
            WARNINGS.append(f"Picture could not be read, skipped: {path.name} ({exc})")
            continue
        pic = fmt(cell.add_paragraph(), before=16, after=8, left=0.1)
        pic.add_run().add_picture(stream, width=Inches(w), height=Inches(h))

    if item.description:
        fmt(add_rich(cell.add_paragraph(), item.description), before=10)

    def section(title, lines, bullet):
        if not lines:
            return
        fmt(add_rich(cell.add_paragraph(), title, bold=True), before=10)
        for text in lines:
            fmt(add_rich(cell.add_paragraph(), ("• " if bullet else "") + text))

    section("Key Features", item.key_features, bullet=True)
    section("Specifications", item.specifications, bullet=True)
    section("In The Box", item.in_the_box, bullet=False)


def keep_together(table):
    """Ask Word not to split a (small) table across two pages."""
    for row in table.rows[:-1]:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.keep_with_next = True


def add_summary_table(doc, items, title, discounted_price):
    """ONE order-summary table: banner (title), item rows, Total, optional Total Discounted Price."""
    total, discounted = totals(items, discounted_price)
    banner_text = package_banner(title)
    first = 2 if banner_text else 1                       # index of the first item row (row 0 = header)
    rows = first + len(items) + 1 + (1 if discounted is not None else 0)
    table = make_table(doc, rows, SUMMARY_COLS)
    center, right = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT
    middle = WD_ALIGN_VERTICAL.CENTER

    # header row
    for cell, label in zip(table.rows[0].cells, ("Line Item", "Quantity", "Description", "Price")):
        cell_text(cell, label, bold=True, align=center, valign=middle)
    row_height(table.rows[0], 0.38)

    # package banner row (merged across all columns) -- skipped when there is no title
    if banner_text:
        banner = table.cell(1, 0).merge(table.cell(1, 3))
        cell_text(banner, banner_text, bold=True, align=center, valign=middle)
        row_height(table.rows[1], 0.27)

    # item rows (numbered 1.0, 2.0 ... inside each table)
    for n, item in enumerate(items, start=1):
        r = table.rows[first - 1 + n].cells
        cell_text(r[0], f"{n}.0", align=center, valign=middle)
        cell_text(r[1], str(item.quantity), align=center, valign=middle)
        desc = cell_text(r[2], item.summary_label, valign=middle)
        if item.optional:
            add_rich(desc, " (Optional)", color="FF0000")
        cell_text(r[3], money(item.total, symbol=False), align=right, valign=middle)
        row_height(table.rows[first - 1 + n], 0.25)

    # Total row, then the optional discounted row
    r = first + len(items)
    cell_text(table.cell(r, 2), "Total", bold=True, align=right, valign=middle)
    cell_text(table.cell(r, 3), money(total), bold=True, align=right, valign=middle)
    row_height(table.rows[r], 0.25)
    if discounted is not None:
        row_height(table.rows[r + 1], 0.25)
        cell_text(table.cell(r + 1, 2), "Total Discounted Price", bold=True, align=right,
                  valign=middle, color="FF0000")
        cell_text(table.cell(r + 1, 3), money(discounted), bold=True, align=right,
                  valign=middle, color="FF0000")
    keep_together(table)
    return table


def add_order_summary(doc, items, discounted_price, extra_summaries=()):
    """Page 3: ORDER SUMMARY. The first table lists ALL items; each extra summary is one more
    table underneath with only the items picked for it (extra = dict: title, items, discounted_price)."""
    rule = fmt(doc.add_paragraph(), after=14)           # horizontal line under the letterhead
    rule.paragraph_format.page_break_before = True      # ORDER SUMMARY always starts a new page
    bottom_border(rule)
    title = fmt(doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=10)
    add_rich(title, "ORDER SUMMARY", bold=True)

    add_summary_table(doc, items, cfg.PACKAGE_TITLE, discounted_price)
    for extra in extra_summaries:
        spacer(doc, 14)
        add_summary_table(doc, extra["items"], extra["title"], extra.get("discounted_price"))
    spacer(doc, 6)


def add_bullet(doc, spec):
    """One bullet line of the terms (spec comes from quotation_config.L())."""
    par = doc.add_paragraph()
    has_marker = bool(spec["marker"])
    fmt(par, before=0, after=spec["after"], left=spec["text_at"],
        first=-(spec["text_at"] - spec["at"]) if has_marker else 0)
    if has_marker:
        par.paragraph_format.tab_stops.add_tab_stop(Inches(spec["text_at"]), WD_TAB_ALIGNMENT.LEFT)
        style_run(par.add_run(spec["marker"] + "\t"))
    add_rich(par, spec["text"])
    return par


def add_timeline_table(doc):
    """Delivery timeline (2 columns; a None lead-time merges with the cell above)."""
    rows = cfg.TIMELINE_ROWS
    table = make_table(doc, 1 + len(rows), TIMELINE_COLS)
    for row in table.rows:
        row_height(row, 0.27)
    center = WD_ALIGN_PARAGRAPH.CENTER
    for cell, label in zip(table.rows[0].cells, cfg.TIMELINE_HEADER):
        cell_text(cell, label, bold=True, align=center, valign=WD_ALIGN_VERTICAL.CENTER)

    for n, (place, lead) in enumerate(rows, start=1):
        cell_text(table.cell(n, 0), place, bold=True, valign=WD_ALIGN_VERTICAL.CENTER)
        if lead is not None:
            cell_text(table.cell(n, 1), lead, bold=True, align=center,
                      valign=WD_ALIGN_VERTICAL.CENTER)
        else:  # merge with the cell above
            table.cell(n - 1, 1).merge(table.cell(n, 1))
    return table


def add_terms(doc):
    """Availability / delivery / payment / validity / warranty section."""
    for spec in cfg.TERMS_BEFORE_TIMELINE:
        heading = add_bullet(doc, spec)
    heading.paragraph_format.keep_with_next = True       # "Delivery Timeline:" never ends a page alone
    keep_together(add_timeline_table(doc))
    spacer(doc, 18)
    for spec in cfg.TERMS_AFTER_TIMELINE:
        add_bullet(doc, spec)
    add_rich(fmt(doc.add_paragraph(), after=14), cfg.NOTE)
    for spec in cfg.TERMS_FINAL:
        add_bullet(doc, spec)


def add_company_block(doc):
    """Bank / company details and the 'Prepared by' signature block."""
    for line in cfg.COMPANY_BLOCK:
        add_rich(fmt(doc.add_paragraph()), line)
    spacer(doc, 28).paragraph_format.keep_with_next = True

    first = fmt(doc.add_paragraph())
    first.paragraph_format.keep_with_next = True
    add_rich(first, f"Prepared By: {cfg.PREPARED_BY[0]}")
    if cfg.SIGNATURE_IMAGE and Path(cfg.SIGNATURE_IMAGE).exists():
        float_picture(first.add_run(), cfg.SIGNATURE_IMAGE, width_in=1.2, x_in=1.1, y_in=-0.5)
    rest = cfg.PREPARED_BY[1:]
    for n, line in enumerate(rest):
        par = add_rich(fmt(doc.add_paragraph(), left=0.97), line)
        par.paragraph_format.keep_with_next = n < len(rest) - 1


def build_document(info, items, discounted_price, extra_summaries=()):
    doc = Document()
    setup_page(doc)
    add_info_box(doc, info)
    add_items_table(doc, items)
    add_order_summary(doc, items, discounted_price, extra_summaries)
    add_terms(doc)
    add_company_block(doc)
    return doc


# ============================================================================
# 5. PDF CONVERSION
# ============================================================================
def find_libreoffice():
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for path in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(path).exists():
            return path
    return None


def _pdf_with_word(docx_path, pdf_path):
    """Windows/Mac: let Microsoft Word do the conversion (needs Word installed)."""
    if sys.platform == "win32":
        import pythoncom                      # needed when running in a background thread
        pythoncom.CoInitialize()
    from docx2pdf import convert
    convert(str(docx_path), str(pdf_path))
    if not Path(pdf_path).exists():
        raise RuntimeError("Word did not produce a PDF")


def _pdf_with_libreoffice(office, docx_path, pdf_path):
    flags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0   # no black console flash
    with tempfile.TemporaryDirectory() as tmp:                # private profile = no lock-ups
        subprocess.run([office, f"-env:UserInstallation=file://{Path(tmp).as_posix()}/profile",
                        "--headless", "--convert-to", "pdf", "--outdir", tmp, str(docx_path)],
                       check=True, capture_output=True, timeout=180, creationflags=flags)
        shutil.move(str(Path(tmp) / (Path(docx_path).stem + ".pdf")), str(pdf_path))


def docx_to_pdf(docx_path, pdf_path):
    """Microsoft Word first (Windows/Mac), then LibreOffice. Raises RuntimeError if neither works."""
    problems = []
    if sys.platform in ("win32", "darwin"):
        try:
            return _pdf_with_word(docx_path, pdf_path)
        except Exception as exc:                              # noqa: BLE001
            problems.append(f"Microsoft Word: {exc}")
    office = find_libreoffice()
    if office:
        try:
            return _pdf_with_libreoffice(office, docx_path, pdf_path)
        except Exception as exc:                              # noqa: BLE001
            problems.append(f"LibreOffice: {exc}")
    else:
        problems.append("LibreOffice not found")
    raise RuntimeError("Could not create the PDF (install Microsoft Word or the free LibreOffice).\n  "
                       + "\n  ".join(problems))


# ============================================================================
# 6. generate()  --  ONE CALL BUILDS THE QUOTATION
# ============================================================================
# Terms / letterhead / bank details live in quotation_config.py. In the .exe they are
# exported once to a JSON file next to the exe so they can be edited without Python.
TEXT_SETTINGS = ["HEADER_LINES", "FOOTER_ADDRESS", "TERMS_BEFORE_TIMELINE", "TIMELINE_HEADER",
                 "TIMELINE_ROWS", "TERMS_AFTER_TIMELINE", "NOTE", "TERMS_FINAL",
                 "COMPANY_BLOCK", "PREPARED_BY"]


def _json_copy(value):
    return json.loads(json.dumps(value))


DEFAULT_WORDING = _json_copy({name: getattr(cfg, name) for name in TEXT_SETTINGS})   # as shipped in quotation_config.py
DEFAULT_PACKAGE_TITLE = cfg.PACKAGE_TITLE


def export_text_settings(path):
    """Write the DEFAULT terms/letterhead/bank text to a JSON file."""
    Path(path).write_text(json.dumps(DEFAULT_WORDING, indent=2, ensure_ascii=False), encoding="utf-8")


def effective_wording(file=None, override=None):
    """The wording a quotation would use: the defaults, with `file` (terms_and_company.json) on top,
    or `override` (wording saved inside a quotation file) instead of the file."""
    wording = _json_copy(DEFAULT_WORDING)
    if override is not None:
        source = override
    elif file and Path(file).exists():
        try:
            source = json.loads(Path(file).read_text(encoding="utf-8"))
        except ValueError as exc:
            raise ValueError(f"{Path(file).name} is not valid JSON (check quotes and commas): {exc}") from exc
    else:
        source = {}
    wording.update({k: v for k, v in source.items() if k in TEXT_SETTINGS})
    return wording


def apply_wording(wording):
    for name, value in wording.items():
        setattr(cfg, name, _json_copy(value))


def package_banner(title):
    """Text of the ORDER SUMMARY banner: 'MATRICE 4 THERMAL' -> 'MATRICE 4 THERMAL PACKAGE'."""
    title = (title or "").strip()
    suffix = getattr(cfg, "PACKAGE_BANNER_SUFFIX", "")
    if title and suffix and not title.upper().endswith(suffix.strip().upper()):
        title += suffix
    return title


DATE_FORMATS = ("%B %d, %Y", "%b %d, %Y", "%B %d %Y", "%b %d %Y", "%d %B %Y", "%d %b %Y", "%Y-%m-%d",
                "%Y/%m/%d", "%m/%d/%Y", "%m-%d-%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y%m%d", "%m/%d/%y")


def date_code(text):
    """'October 2, 2026' -> '20261002'.  Unreadable dates fall back to today's date."""
    cleaned = " ".join((text or "").replace(".", " ").replace(",", ", ").split()).replace(" ,", ",")
    for fmt_ in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(cleaned, fmt_).strftime("%Y%m%d")
        except ValueError:
            continue
    return datetime.date.today().strftime("%Y%m%d")


def output_stem(info, package_title):
    """File name (without extension): Q-26-626 - 20261002 - MATRICE 4 THERMAL - PNP Mindoro"""
    values = {
        "number": info.get("quotation_number", ""),
        "date": date_code(info.get("date", "")),
        "package": package_title or "",
        "client": info.get("company_name", "") or info.get("attention_to", ""),
    }
    parts = [safe_filename(values[name]) for name in cfg.FILENAME_PARTS if safe_filename(values[name])]
    return (cfg.FILENAME_SEPARATOR.join(parts) or "Quotation")[:150].rstrip(" .")


def safe_filename(text):
    """Make text safe for a Windows file name: drop " ? * < >, turn / \\ : | into '-'."""
    out = []
    for ch in str(text or ""):
        if ch in '"?*<>' or ord(ch) < 32:
            continue
        out.append("-" if ch in "/\\:|" else ch)
    return " ".join("".join(out).split()).strip(" .")


# ----------------------------------------------------------------------------
# project file (.qgen): everything needed to re-open and edit a quotation later
# ----------------------------------------------------------------------------
PROJECT_EXT = ".qgen"
PROJECT_FORMAT = 1


def write_project_file(path, payload, items, images_dir):
    """Save `payload` as JSON, plus a compressed copy of every picture used (so the file still
    opens correctly on another PC / after the picture was deleted)."""
    pictures = {}
    for item in items:
        for name in item.images:
            src = Path(images_dir) / name
            if name not in pictures and src.exists():
                try:
                    stream, _, _ = prepare_picture(src)
                    pictures[name] = base64.b64encode(stream.read()).decode("ascii")
                except Exception:                                 # noqa: BLE001
                    pass
    path = Path(path)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps({**payload, "pictures": pictures}, indent=1, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, path)


def read_project_file(path):
    """Load a .qgen file (raises ValueError if it is not one)."""
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except ValueError as exc:
        raise ValueError("This is not a valid quotation file.") from exc
    if not isinstance(data, dict) or data.get("format") != PROJECT_FORMAT or "items" not in data:
        raise ValueError("This is not a quotation file made by Keith's Quotation Generator.")
    return data


def restore_project_pictures(data, images_dir):
    """Re-create pictures that are missing from the pictures folder. Returns how many were restored."""
    restored = 0
    for name, encoded in (data.get("pictures") or {}).items():
        if not name or any(sep in name for sep in "/\\"):
            continue                                              # only plain file names are restored
        dest = Path(images_dir) / name
        if not dest.exists():
            try:
                dest.parent.mkdir(parents=True, exist_ok=True)
                dest.write_bytes(base64.b64decode(encoded))
                restored += 1
            except (OSError, ValueError):
                pass
    return restored


# ============================================================================
# 6b. generate()  --  ONE CALL BUILDS THE QUOTATION
# ============================================================================
def resolve_extra_summaries(items, extra_summaries):
    """[{title, uids, discounted_price}] -> [{title, items, discounted_price}] (quotation order kept)."""
    resolved = []
    for n, extra in enumerate(extra_summaries or [], start=2):
        wanted = set(extra.get("uids", []))
        chosen = [i for i in items if i.uid in wanted]
        if not chosen:
            raise ValueError(f"Order summary #{n} ('{extra.get('title', '')}') has no items. "
                             f"Edit it or remove it.")
        resolved.append({"title": extra.get("title", ""), "items": chosen,
                         "discounted_price": extra.get("discounted_price")})
    return resolved


def generate(info, items, out_dir, *, discounted_price=None, extra_summaries=None, images_dir=None,
             package_title=None, logo=None, signature=None, make_docx=True, make_pdf=True,
             save_project=True, text_settings_file=None, text_settings=None):
    """Build the quotation: Word, PDF and the editable .qgen file.
    Returns {"docx", "pdf", "project": Path|None, "pdf_error", "warnings", "items", "total", "stem"}."""
    WARNINGS.clear()
    wording = effective_wording(text_settings_file, text_settings)
    apply_wording(wording)
    if images_dir:
        cfg.IMAGES_DIR = Path(images_dir)
    cfg.PACKAGE_TITLE = package_title if package_title is not None else DEFAULT_PACKAGE_TITLE
    if logo is not None:
        cfg.LOGO_IMAGE = Path(logo) if logo else None
    if signature is not None:
        cfg.SIGNATURE_IMAGE = Path(signature) if signature else None

    blank = {"quotation_number": "", "date": "", "company_name": "", "attention_to": "",
             "contact": "", "email": "", "address": ""}
    info = {**blank, **info}                                  # any field not supplied stays blank
    if not info.get("date"):
        today = datetime.date.today()
        info["date"] = f"{today:%B} {today.day}, {today.year}"

    if not items:
        raise ValueError("The quotation has no items.")
    extras = resolve_extra_summaries(items, extra_summaries)
    doc = build_document(info, items, discounted_price, extras)

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = output_stem(info, cfg.PACKAGE_TITLE)
    docx_path, pdf_path = out_dir / f"{stem}.docx", out_dir / f"{stem}.pdf"
    project_path = out_dir / f"{stem}{PROJECT_EXT}"
    result = {"docx": None, "pdf": None, "project": None, "pdf_error": None, "items": items,
              "total": totals(items, discounted_price)[0], "stem": stem}

    if make_docx:
        doc.save(docx_path)
        result["docx"] = docx_path
    if make_pdf:
        try:
            if make_docx:
                docx_to_pdf(docx_path, pdf_path)
            else:                                             # PDF only: keep Word file in a temp folder
                with tempfile.TemporaryDirectory() as tmp:
                    tmp_docx = Path(tmp) / f"{stem}.docx"
                    doc.save(tmp_docx)
                    docx_to_pdf(tmp_docx, pdf_path)
            result["pdf"] = pdf_path
        except Exception as exc:                              # noqa: BLE001
            result["pdf_error"] = str(exc)

    if save_project:                                          # written last, so it can record what was generated
        tables = [{"title": package_banner(cfg.PACKAGE_TITLE), "items": len(items),
                   "total": totals(items, discounted_price)[0], "discounted_price": discounted_price}]
        tables += [{"title": package_banner(e["title"]), "items": len(e["items"]),
                    "total": totals(e["items"], None)[0], "discounted_price": e["discounted_price"]}
                   for e in extras]
        payload = {
            "format": PROJECT_FORMAT, "app": "Keith's Quotation Generator",
            "saved_at": datetime.datetime.now().isoformat(timespec="seconds"),
            "info": info, "package_title": cfg.PACKAGE_TITLE, "discounted_price": discounted_price,
            "items": [item_to_dict(i) for i in items],
            "extra_summaries": [{"title": e.get("title", ""), "uids": list(e.get("uids", [])),
                                 "discounted_price": e.get("discounted_price")}
                                for e in (extra_summaries or [])],
            "wording": wording,
            "generated": {"stem": stem, "order_summary_tables": tables,
                          "files": {"docx": docx_path.name if result["docx"] else None,
                                    "pdf": pdf_path.name if result["pdf"] else None}},
        }
        write_project_file(project_path, payload, items, cfg.IMAGES_DIR)
        result["project"] = project_path
    result["warnings"] = list(WARNINGS)
    return result


# ============================================================================
# 7. COMMAND LINE  (optional -- the GUI in quotation_app.py is the normal way to use this)
# ============================================================================
def parse_args():
    p = argparse.ArgumentParser(description="Generate a quotation (Word + PDF). "
                                "Any option given here overrides quotation_config.py.")
    p.add_argument("--csv", help="item list; ALL rows are put in the quotation (default: products.csv)")
    p.add_argument("--number", help="Quotation Number")
    p.add_argument("--date", help="Date, e.g. 'October 2, 2026' (default: config, or today)")
    p.add_argument("--company", help="Company Name")
    p.add_argument("--attention", help="Attention To")
    p.add_argument("--contact", help="Contact")
    p.add_argument("--email", help="Email")
    p.add_argument("--address", help="Address")
    p.add_argument("--discounted-price", type=float,
                   help="show a 'Total Discounted Price' row with this amount (default: none)")
    p.add_argument("--no-pdf", action="store_true", help="create the Word file only")
    return p.parse_args()


def main():
    args = parse_args()
    info = dict(cfg.QUOTE_INFO)
    for key, value in (("quotation_number", args.number), ("date", args.date),
                       ("company_name", args.company), ("attention_to", args.attention),
                       ("contact", args.contact), ("email", args.email), ("address", args.address)):
        if value is not None:
            info[key] = value
    discounted = args.discounted_price if args.discounted_price is not None else cfg.DISCOUNTED_PRICE
    csv_path = Path(args.csv or cfg.PRODUCTS_CSV)
    result = generate(info, load_items(csv_path), cfg.OUTPUT_DIR, images_dir=csv_path.parent / "images",
                      discounted_price=discounted, make_pdf=not args.no_pdf)
    for w in result["warnings"]:
        print("!", w)
    print("Word :", result["docx"])
    print("PDF  :", result["pdf"] or f"(failed) {result['pdf_error']}")


if __name__ == "__main__":
    main()
